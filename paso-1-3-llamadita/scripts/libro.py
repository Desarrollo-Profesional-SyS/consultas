# Arma el libro "Paso 1.3" con la identidad de academy.dpsys.com.mx.
# Las hojas de resumen son formulas sobre las dos hojas de datos.
import os, sys
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.comments import Comment

D, SALIDA, LOGO = sys.argv[1], sys.argv[2], sys.argv[3]
TZ = ZoneInfo('America/Matamoros')
DESDE = datetime(2026, 10, 5, tzinfo=TZ)
CORTE_TXT = '27/10/2026 15:00 (hora de Reynosa)'

# Identidad (variables de sys-site.css)
MAGENTA, ROSA, INK, VINO = '800E64', 'E91E63', '330A28', '880E4F'
SUPERFICIE, TINTE, ROSA_TINTE, LINEA, GRIS, NARANJA = 'FBF8FA', 'F6EBF3', 'FDEDF3', 'E8DEE6', '6B6B6B', 'FF9800'
CUERPO, DISPLAY = 'Poppins', 'Anton'

def fnt(size=10, bold=False, color=INK, name=CUERPO, italic=False):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)
def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)
fino = Side(style='thin', color=LINEA)
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
RAYA_ROSA = Border(bottom=Side(style='medium', color=ROSA))

P = pd.read_pickle(os.path.join(D, 'P.pkl'))
E = pd.read_pickle(os.path.join(D, 'E2.pkl'))
P = P[P.t_pregunta >= DESDE].copy()
E = E[E.t_13 >= DESDE].copy()

def sn(v):
    return 'Sí' if bool(v) else 'No'
def num(v, nd=1):
    return None if v is None or pd.isna(v) else round(float(v), nd)
def semana(t):
    l = t.astimezone(TZ)
    return (l - timedelta(days=l.weekday())).date()
def franja(t):
    h = t.astimezone(TZ).hour
    return '9 a 18 h' if 9 <= h < 18 else ('18 a 22 h' if 18 <= h < 22 else '22 a 9 h')

SALIDAS = {'NEG': 'Negativo', 'S14': '1.4 Facilidades de pago', 'PAGO': 'Medio de pago', 'GAN': 'Ganado', 'PER': 'Perdido',
           'SIGUE_EN_13': 'Sigue en 1.3', 'S12': '1.2 Lista de cursos', 'S11': '1.1 Informes', 'S1': '1 Saludo', 'STBY': 'STBY', 'ENT': 'Leads entrantes', 'OTRA': 'Otra'}

def no_movido(r):
    if r['movido_13']: return ''
    k, o = r.get('salida_12'), r.get('salida_12_origen')
    if k is None or pd.isna(k): return 'Sigue en 1.2'
    if k == 'NEG': return 'Negativo (lo mandó el bot)' if o == 'secuencia' else 'Negativo (lo movió una persona)'
    if k in ('S14', 'PAGO', 'GAN', 'STBY'): return '1.4, pago o Ganado (persona)'
    if k == 'PER': return 'Perdido (persona)'
    return SALIDAS.get(k, k)

# ---------------------------------------------------------------- datos
filas_base = []
for _, r in P.sort_values('t_pregunta').iterrows():
    contesto = pd.notna(r.t_resp)
    a_tiempo = bool(r.contesto_en_12)
    cuando = 'No contestó' if not contesto else ('A tiempo (en 1.2)' if a_tiempo else 'Tarde (ya fuera de 1.2)')
    f = [int(r.lead), r.embudo, semana(r.t_pregunta).strftime('%d/%m/%Y'), sn(contesto), cuando]
    if contesto:
        f += [num(r.min_resp), sn(r.pregunta_precio), r.humano_quien or 'Nadie', num(r.min_humano), sn(r.plantilla_audio), sn(r.audio),
              r.audio_quien or '', num(r.min_audio), sn(r.movido_13), r.quien_13 if r.movido_13 else '', num(r.min_13) if r.movido_13 else None,
              sn(r['13_antes_neg']) if r.movido_13 else '', r.orden, no_movido(r), franja(r.t_resp)]
    else:
        f += [None] * 15
    f += [sn(r.llego_14_o_mas), sn(r.llego_pago), sn(r.venta), sn(r.maduro), r.etapa_actual or '']
    filas_base.append(f)
CAB_BASE = ['Lead', 'Embudo', 'Semana (lunes)', 'Contestó la pregunta del área', 'Cuándo contestó', 'Min. en contestar', 'Preguntó precio o pago',
            'Primera persona que le contestó', 'Min. a la primera respuesta de una persona', 'Plantilla "audio rapidito"', 'Le mandaron audio', 'Quién mandó el audio',
            'Min. al audio', 'Movido a 1.3', 'Quién lo movió a 1.3', 'Min. de su respuesta a 1.3', 'Movido antes de que el bot lo mandara al Negativo',
            'Qué hicieron después de su respuesta', 'Si no se movió a 1.3, a dónde fue', 'Franja de su respuesta', 'Llegó a 1.4 o más', 'Llegó a un medio de pago',
            'Venta (Ganado o pago confirmado)', 'Con 7 días o más', 'Etapa actual']

filas_e = []
for _, r in E.sort_values('t_13').iterrows():
    sig = SALIDAS.get(r.siguiente, r.siguiente)
    saco = '' if r.siguiente == 'SIGUE_EN_13' else ('Bot de 1.3' if r.siguiente_origen == 'secuencia' else (r.siguiente_quien or 'Persona'))
    filas_e.append([int(r.lead), r.embudo, semana(r.t_13).strftime('%d/%m/%Y'), r.t_13.astimezone(TZ).strftime('%d/%m/%Y %H:%M'), SALIDAS.get(r.desde_k, r.desde or ''), r.quien,
                    sn(r.bot_llamadita), sn(r.bot_facilidades), sn(r.cliente_respondio), num(r.min_resp_cliente), sn(r.humano_escribio), sn(r.tarea_llamada),
                    sig, saco, num(r.horas_en_13), sn(r.despues_14_o_mas), sn(r.despues_pago), sn(r.despues_gano), sn(r.cliente_resp_y_bot_a_neg), r.asignado or '',
                    '' if r.resp30_llamadita is None else sn(r.resp30_llamadita), '' if r.resp30_facilidades is None else sn(r.resp30_facilidades)])
CAB_E = ['Lead', 'Embudo', 'Semana (lunes)', 'Entró a 1.3', 'Venía de', 'Quién lo movió', 'Bot mandó "llamadita"', 'Bot mandó "facilidades de pago"',
         'El cliente respondió en 1.3', 'Min. a su respuesta', 'Una persona le escribió en 1.3', 'Tarea de llamada', 'Siguiente etapa', 'Quién lo sacó de 1.3',
         'Horas en 1.3', 'Después llegó a 1.4 o más', 'Después llegó a un medio de pago', 'Después Ganado', 'Respondió y el bot lo mandó al Negativo', 'Asesor asignado',
         'Respondió en 30 min a la llamadita', 'Respondió en 30 min a facilidades']

NB, NE = len(filas_base), len(filas_e)
B_FIN, E_FIN = NB + 4, NE + 4  # datos desde la fila 5

def rng(hoja, col, fin):
    return f"'{hoja}'!${col}$5:${col}${fin}"
def bcol(c):
    return rng('Base leads', c, B_FIN)
def ecol(c):
    return rng('Entradas a 1.3', c, E_FIN)

wb = Workbook()

def encabezado(ws, titulo, subtitulo, ancho_cols):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(ancho_cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 46
    ws.row_dimensions[2].height = 20
    ws['B1'] = titulo
    ws['B1'].font = Font(name=DISPLAY, size=22, color=MAGENTA)
    ws['B1'].alignment = Alignment(vertical='center')
    ws['B2'] = subtitulo
    ws['B2'].font = fnt(10, color=GRIS)
    for c in range(1, len(ancho_cols) + 1):
        ws.cell(row=3, column=c).border = RAYA_ROSA
    ws.row_dimensions[3].height = 6

def seccion(ws, fila, texto, ncols=8):
    c = ws.cell(row=fila, column=2, value=texto.upper())
    c.font = Font(name=DISPLAY, size=13, color=MAGENTA)
    ws.row_dimensions[fila].height = 24
    for col in range(2, 2 + ncols):
        ws.cell(row=fila, column=col).border = Border(bottom=Side(style='thin', color=ROSA))

def cab_tabla(ws, fila, textos, col0=2):
    ws.row_dimensions[fila].height = 32
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + i, value=t)
        c.font = fnt(9, True, 'FFFFFF'); c.fill = fill(MAGENTA)
        c.alignment = Alignment(horizontal='center' if i else 'left', vertical='center', wrap_text=True)
        c.border = BORDE

def celda(ws, fila, col, valor, fmt=None, bold=False, color=INK, banda=False, alin=None):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = fnt(10, bold, color)
    if fmt: c.number_format = fmt
    if banda: c.fill = fill(SUPERFICIE)
    c.border = BORDE
    c.alignment = Alignment(horizontal=alin or ('left' if col == 2 else 'center'), vertical='center', wrap_text=(col == 2))
    return c

def nota(ws, fila, texto, col=2):
    c = ws.cell(row=fila, column=col, value=texto)
    c.font = fnt(9, color=GRIS, italic=True)
    c.alignment = Alignment(wrap_text=False, vertical='top')

PCT, ENT, DEC = '0.0%;(0.0%);"-"', '#,##0;(#,##0);"-"', '0.0;(0.0);"-"'
def div(a, b):
    return f'=IF({b}=0,0,{a}/{b})'
def mediana(rango, conds):
    cond = '*'.join(f'({c})' for c in conds)
    return f'MEDIAN(IF({cond},IF({rango}<>"",{rango})))'

EMB = ['ISO + Calidad', 'Core Tools Premium', 'Manufactura y Calidad']

# ================================================================ Resumen
ws = wb.active; ws.title = 'Resumen'
encabezado(ws, 'Paso 1.3 · Llamadita y facilidades de pago', f'ISO + Calidad, Core Tools Premium y Manufactura y Calidad · leads del 5/10/2026 al corte del {CORTE_TXT} · lectura del CRM sin cambios',
           [3, 62, 16, 16, 16, 16, 3])
fila = 5
seccion(ws, fila, 'Lo que se encontró', 5); fila += 1
HALLAZGOS_FILA = fila
fila += 14  # espacio para hallazgos y recomendaciones
seccion(ws, fila, 'Cifras clave', 5); fila += 1
cab_tabla(ws, fila, ['Indicador', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); fila += 1
KPI = {}
def kpi(nombre, f_total, f_emb, fmt, clave=None):
    global fila
    banda = (fila % 2 == 0)
    celda(ws, fila, 2, nombre, banda=banda)
    celda(ws, fila, 3, f_total, fmt, True, MAGENTA, banda)
    for i, e in enumerate(EMB):
        v = f_emb(e) if f_emb else None
        celda(ws, fila, 4 + i, v, fmt, banda=banda)
    if clave: KPI[clave] = fila
    fila += 1

def cnt(*pares, hoja='b'):
    col = bcol if hoja == 'b' else ecol
    return 'COUNTIFS(' + ','.join(f'{col(c)},"{v}"' for c, v in pares) + ')'
def cnt_e(emb, *pares, hoja='b'):
    return cnt(('B', emb), *pares, hoja=hoja)

kpi('Leads que recibieron la pregunta "¿en qué área trabajas?"', f'=COUNTA({bcol("A")})', lambda e: f'={cnt_e(e)}', ENT, 'preg')
kpi('Contestaron a tiempo (antes de que el bot los mandara al Negativo)', f'={cnt(("E", "A tiempo (en 1.2)"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"))}', ENT, 'resp')
r = KPI['resp']; q = KPI['preg']
kpi('% que contesta a tiempo', div(f'C{r}', f'C{q}'), lambda e: None, PCT, 'resp_pct')
for i in range(3):
    col = get_column_letter(4 + i)
    ws.cell(row=KPI['resp_pct'], column=4 + i, value=div(f'{col}{r}', f'{col}{q}')).number_format = PCT
kpi('De los que contestaron a tiempo: les mandaron audio', f'={cnt(("E", "A tiempo (en 1.2)"), ("K", "Sí"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("K", "Sí"))}', ENT, 'audio')
kpi('De los que contestaron a tiempo: movidos a 1.3', f'={cnt(("E", "A tiempo (en 1.2)"), ("N", "Sí"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("N", "Sí"))}', ENT, 'mov')
kpi('% movido a 1.3 (de los que contestaron a tiempo)', None, None, PCT, 'mov_pct')
for i in range(4):
    col = get_column_letter(3 + i)
    c = ws.cell(row=KPI['mov_pct'], column=3 + i, value=div(f'{col}{KPI["mov"]}', f'{col}{r}')); c.number_format = PCT
    if i == 0: c.font = fnt(10, True, MAGENTA)
kpi('Contestaron a tiempo, recibieron audio y NO los movieron a 1.3', f'={cnt(("E", "A tiempo (en 1.2)"), ("R", "Audio sin mover a 1.3"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("R", "Audio sin mover a 1.3"))}', ENT, 'audio_sin')
kpi('Contestaron a tiempo y no recibieron ni audio ni 1.3', f'={cnt(("E", "A tiempo (en 1.2)"), ("R", "Ni audio ni 1.3"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("R", "Ni audio ni 1.3"))}', ENT, 'nada')
kpi('Contestaron a tiempo, no los movieron y el bot los mandó al Negativo', f'={cnt(("E", "A tiempo (en 1.2)"), ("S", "Negativo (lo mandó el bot)"))}', lambda e: f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("S", "Negativo (lo mandó el bot)"))}', ENT, 'neg_bot')
kpi('Minutos de su respuesta a 1.3 (mediana, movidos)', ArrayFormula(f'C{fila}', '=' + mediana(bcol('P'), [f'{bcol("E")}="A tiempo (en 1.2)"', f'{bcol("N")}="Sí"'])), None, DEC, 'min13')
for i, e in enumerate(EMB):
    ref = f'{get_column_letter(4 + i)}{KPI["min13"]}'
    ws[ref] = ArrayFormula(ref, '=' + mediana(bcol('P'), [f'{bcol("E")}="A tiempo (en 1.2)"', f'{bcol("N")}="Sí"', f'{bcol("B")}="{e}"']))
    ws[ref].number_format = DEC; ws[ref].font = fnt(10); ws[ref].border = BORDE; ws[ref].alignment = Alignment(horizontal='center')
kpi('Entradas a 1.3 en el periodo', f'=COUNTA({ecol("A")})', lambda e: f'={cnt_e(e, hoja="e")}', ENT, 'e13')
kpi('Entradas a 1.3 en las que el cliente respondió', f'={cnt(("I", "Sí"), hoja="e")}', lambda e: f'={cnt_e(e, ("I", "Sí"), hoja="e")}', ENT, 'e13_resp')
def tasa(num_par, den_par, emb=None):
    n = cnt_e(emb, num_par, hoja='e') if emb else cnt(num_par, hoja='e')
    d = cnt_e(emb, den_par, hoja='e') if emb else cnt(den_par, hoja='e')
    return f'=IF({d}=0,0,{n}/{d})'
kpi('Respuesta en 30 min a la "llamadita" del bot', tasa(('U', 'Sí'), ('G', 'Sí')), lambda e: tasa(('U', 'Sí'), ('G', 'Sí'), e), PCT, 'r30l')
kpi('Respuesta en 30 min a "facilidades de pago" del bot', tasa(('V', 'Sí'), ('H', 'Sí')), lambda e: tasa(('V', 'Sí'), ('H', 'Sí'), e), PCT, 'r30f')
kpi('Entradas a 1.3 que después llegaron a 1.4 o más', f'={cnt(("P", "Sí"), hoja="e")}', lambda e: f'={cnt_e(e, ("P", "Sí"), hoja="e")}', ENT, 'e13_14')
kpi('% de entradas a 1.3 que llega a 1.4 o más', None, None, PCT, 'e13_14_pct')
for i in range(4):
    col = get_column_letter(3 + i)
    c = ws.cell(row=KPI['e13_14_pct'], column=3 + i, value=div(f'{col}{KPI["e13_14"]}', f'{col}{KPI["e13"]}')); c.number_format = PCT
    if i == 0: c.font = fnt(10, True, MAGENTA)
kpi('Entradas a 1.3 que terminaron en Ganado', f'={cnt(("R", "Sí"), hoja="e")}', lambda e: f'={cnt_e(e, ("R", "Sí"), hoja="e")}', ENT, 'e13_gan')
kpi('Salieron de 1.3 al Negativo por el bot', f'={cnt(("M", "Negativo"), ("N", "Bot de 1.3"), hoja="e")}', lambda e: f'={cnt_e(e, ("M", "Negativo"), ("N", "Bot de 1.3"), hoja="e")}', ENT, 'e13_neg')
kpi('El cliente respondió en 1.3 y aun así el bot lo mandó al Negativo', f'={cnt(("S", "Sí"), hoja="e")}', lambda e: f'={cnt_e(e, ("S", "Sí"), hoja="e")}', ENT, 'e13_resp_neg')
kpi('Entradas a 1.3 con tarea de llamada creada', f'={cnt(("L", "Sí"), hoja="e")}', lambda e: f'={cnt_e(e, ("L", "Sí"), hoja="e")}', ENT, 'e13_tarea')
for rr in range(KPI['preg'], fila):
    for cc in range(2, 7):
        ws.cell(row=rr, column=cc).border = BORDE
fila += 1
nota(ws, fila, 'Cada cifra es una fórmula sobre las hojas "Base leads" y "Entradas a 1.3". El método y las definiciones están en "Fuentes y método".')
RESUMEN_FIN = fila

# ================================================================ Qué pasa tras el área
w2 = wb.create_sheet('Tras la pregunta del área')
encabezado(w2, 'Qué pasa cuando el lead contesta el área', 'Leads que contestaron "¿en qué área trabajas?" antes de que el bot los mandara al Negativo (2 h)', [3, 50, 15, 15, 15, 15, 15, 3])
f2 = 5
seccion(w2, f2, 'Lo que hizo el equipo después de su respuesta', 5); f2 += 1
cab_tabla(w2, f2, ['Qué hicieron', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura', '% del total']); f2 += 1
ORDENES = ['Audio y después 1.3', '1.3 y después audio', 'Movido a 1.3 sin audio', 'Audio sin mover a 1.3', 'Ni audio ni 1.3']
ini = f2
for o in ORDENES:
    b = f2 % 2 == 0
    celda(w2, f2, 2, o, banda=b)
    celda(w2, f2, 3, f'={cnt(("E", "A tiempo (en 1.2)"), ("R", o))}', ENT, True, MAGENTA, b)
    for i, e in enumerate(EMB):
        celda(w2, f2, 4 + i, f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("R", o))}', ENT, banda=b)
    celda(w2, f2, 7, div(f'C{f2}', f'C${ini + len(ORDENES)}'), PCT, banda=b)
    f2 += 1
celda(w2, f2, 2, 'Total que contestó a tiempo', bold=True)
for c in range(3, 7):
    L = get_column_letter(c)
    celda(w2, f2, c, f'=SUM({L}{ini}:{L}{f2 - 1})', ENT, True, MAGENTA if c == 3 else INK)
celda(w2, f2, 7, div(f'C{f2}', f'C{f2}'), PCT, True)
TOT2 = f2
f2 += 2
nota(w2, f2, 'El protocolo esperado es: plantilla "Excelente… te mando un audio rapidito", nota de voz con la recomendación y mover a mano a 1.3.'); f2 += 2

seccion(w2, f2, 'Tiempos (mediana en minutos desde su respuesta)', 5); f2 += 1
cab_tabla(w2, f2, ['Paso', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); f2 += 1
for txt, col, extra in [('Primera respuesta de una persona', 'I', None), ('Audio', 'M', ('K', 'Sí')), ('Movido a 1.3', 'P', ('N', 'Sí'))]:
    b = f2 % 2 == 0
    celda(w2, f2, 2, txt, banda=b)
    base_c = [f'{bcol("E")}="A tiempo (en 1.2)"'] + ([f'{bcol(extra[0])}="{extra[1]}"'] if extra else [])
    for i, e in enumerate([None] + EMB):
        ref = f'{get_column_letter(3 + i)}{f2}'
        conds = base_c + ([f'{bcol("B")}="{e}"'] if e else [])
        w2[ref] = ArrayFormula(ref, '=' + mediana(bcol(col), conds))
        w2[ref].number_format = DEC; w2[ref].font = fnt(10, i == 0, MAGENTA if i == 0 else INK); w2[ref].border = BORDE
        w2[ref].alignment = Alignment(horizontal='center')
        if b: w2[ref].fill = fill(SUPERFICIE)
    f2 += 1
f2 += 1
seccion(w2, f2, 'Los que no movieron a 1.3: a dónde fueron', 5); f2 += 1
cab_tabla(w2, f2, ['Destino', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura', 'Preguntaron precio o pago']); f2 += 1
for d in ['Negativo (lo mandó el bot)', 'Perdido (persona)', '1.4, pago o Ganado (persona)', 'Negativo (lo movió una persona)', 'Sigue en 1.2']:
    b = f2 % 2 == 0
    celda(w2, f2, 2, d, banda=b)
    celda(w2, f2, 3, f'={cnt(("E", "A tiempo (en 1.2)"), ("S", d))}', ENT, True, MAGENTA, b)
    for i, e in enumerate(EMB):
        celda(w2, f2, 4 + i, f'={cnt_e(e, ("E", "A tiempo (en 1.2)"), ("S", d))}', ENT, banda=b)
    celda(w2, f2, 7, f'={cnt(("E", "A tiempo (en 1.2)"), ("S", d), ("G", "Sí"))}', ENT, banda=b)
    f2 += 1
f2 += 1
seccion(w2, f2, 'Según la hora en que contestó (hora de Reynosa)', 5); f2 += 1
cab_tabla(w2, f2, ['Franja', 'Contestaron a tiempo', 'Con audio', 'Movidos a 1.3', '% movidos', 'Min. a 1.3 (mediana)']); f2 += 1
for fr in ['9 a 18 h', '18 a 22 h', '22 a 9 h']:
    b = f2 % 2 == 0
    celda(w2, f2, 2, fr, banda=b)
    celda(w2, f2, 3, f'={cnt(("E", "A tiempo (en 1.2)"), ("T", fr))}', ENT, banda=b)
    celda(w2, f2, 4, f'={cnt(("E", "A tiempo (en 1.2)"), ("T", fr), ("K", "Sí"))}', ENT, banda=b)
    celda(w2, f2, 5, f'={cnt(("E", "A tiempo (en 1.2)"), ("T", fr), ("N", "Sí"))}', ENT, banda=b)
    celda(w2, f2, 6, div(f'E{f2}', f'C{f2}'), PCT, True, MAGENTA, b)
    ref = f'G{f2}'
    w2[ref] = ArrayFormula(ref, '=' + mediana(bcol('P'), [f'{bcol("E")}="A tiempo (en 1.2)"', f'{bcol("N")}="Sí"', f'{bcol("T")}="{fr}"']))
    w2[ref].number_format = DEC; w2[ref].font = fnt(10); w2[ref].border = BORDE; w2[ref].alignment = Alignment(horizontal='center')
    if b: w2[ref].fill = fill(SUPERFICIE)
    f2 += 1
f2 += 1
seccion(w2, f2, 'Los que contestaron tarde (ya estaban en el Negativo)', 5); f2 += 1
cab_tabla(w2, f2, ['Indicador', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); f2 += 1
for txt, pares in [('Contestaron tarde', [('E', 'Tarde (ya fuera de 1.2)')]), ('Les mandaron audio', [('E', 'Tarde (ya fuera de 1.2)'), ('K', 'Sí')]),
                   ('Los movieron a 1.3', [('E', 'Tarde (ya fuera de 1.2)'), ('N', 'Sí')])]:
    b = f2 % 2 == 0
    celda(w2, f2, 2, txt, banda=b)
    celda(w2, f2, 3, f'={cnt(*pares)}', ENT, True, MAGENTA, b)
    for i, e in enumerate(EMB):
        celda(w2, f2, 4 + i, f'={cnt_e(e, *pares)}', ENT, banda=b)
    f2 += 1

# ================================================================ Quién lo usa
w3 = wb.create_sheet('Quién lo usa')
encabezado(w3, 'Quién hace el paso a 1.3', 'Por la primera persona del equipo que le contestó al lead después de su respuesta al área (contestaron a tiempo)', [3, 34, 14, 14, 14, 14, 14, 14, 14, 3])
f3 = 5
seccion(w3, f3, 'Por persona que atendió la respuesta', 8); f3 += 1
cab_tabla(w3, f3, ['Persona', 'Respuestas que atendió', 'Usó la plantilla del audio', 'Mandó audio (cualquiera)', 'Movidos a 1.3', '% con audio', '% movidos a 1.3', 'Min. a la 1ª respuesta (mediana)']); f3 += 1
personas = [p for p in P[P.contesto_en_12].humano_quien.fillna('Nadie').value_counts().index]
ini3 = f3
FILA_PERSONA = {}
for p in personas:
    b = f3 % 2 == 0
    FILA_PERSONA[p] = f3
    celda(w3, f3, 2, p, banda=b)
    celda(w3, f3, 3, f'={cnt(("E", "A tiempo (en 1.2)"), ("H", p))}', ENT, banda=b)
    celda(w3, f3, 4, f'={cnt(("E", "A tiempo (en 1.2)"), ("H", p), ("J", "Sí"))}', ENT, banda=b)
    celda(w3, f3, 5, f'={cnt(("E", "A tiempo (en 1.2)"), ("H", p), ("K", "Sí"))}', ENT, banda=b)
    celda(w3, f3, 6, f'={cnt(("E", "A tiempo (en 1.2)"), ("H", p), ("N", "Sí"))}', ENT, banda=b)
    celda(w3, f3, 7, div(f'E{f3}', f'C{f3}'), PCT, banda=b)
    celda(w3, f3, 8, div(f'F{f3}', f'C{f3}'), PCT, True, MAGENTA, b)
    ref = f'I{f3}'
    w3[ref] = ArrayFormula(ref, '=' + mediana(bcol('I'), [f'{bcol("E")}="A tiempo (en 1.2)"', f'{bcol("H")}="{p}"']))
    w3[ref].number_format = DEC; w3[ref].font = fnt(10); w3[ref].border = BORDE; w3[ref].alignment = Alignment(horizontal='center')
    if b: w3[ref].fill = fill(SUPERFICIE)
    f3 += 1
f3 += 1
nota(w3, f3, 'El CRM le permite a la auxiliar mover etapas: en el mismo periodo movió leads a Perdido, a 1.1 Informes y a Ganado, pero ninguno a 1.3.'); f3 += 2
seccion(w3, f3, 'Quién mueve a 1.3 (todas las entradas a 1.3)', 8); f3 += 1
cab_tabla(w3, f3, ['Persona', 'Entradas a 1.3', 'Desde 1.2', 'Desde el Negativo', 'Llegaron a 1.4 o más', '% a 1.4 o más', 'Ganado']); f3 += 1
for p in E.quien.value_counts().index:
    b = f3 % 2 == 0
    celda(w3, f3, 2, p, banda=b)
    celda(w3, f3, 3, f'={cnt(("F", p), hoja="e")}', ENT, banda=b)
    celda(w3, f3, 4, f'={cnt(("F", p), ("E", "1.2 Lista de cursos"), hoja="e")}', ENT, banda=b)
    celda(w3, f3, 5, f'={cnt(("F", p), ("E", "Negativo"), hoja="e")}', ENT, banda=b)
    celda(w3, f3, 6, f'={cnt(("F", p), ("P", "Sí"), hoja="e")}', ENT, banda=b)
    celda(w3, f3, 7, div(f'F{f3}', f'C{f3}'), PCT, True, MAGENTA, b)
    celda(w3, f3, 8, f'={cnt(("F", p), ("R", "Sí"), hoja="e")}', ENT, banda=b)
    f3 += 1
f3 += 1
nota(w3, f3, 'Todas las entradas a 1.3 del periodo fueron movimientos manuales: ningún bot ni regla mueve al lead a 1.3.')

# ================================================================ Frecuencia
w4 = wb.create_sheet('Frecuencia')
encabezado(w4, 'Con qué frecuencia se usa 1.3', 'Entradas a 1.3 por semana (lunes a domingo) y embudo', [3, 22, 15, 15, 15, 15, 18, 18, 3])
f4 = 5
seccion(w4, f4, 'Entradas a 1.3 por semana', 7); f4 += 1
cab_tabla(w4, f4, ['Semana (lunes)', 'ISO + Calidad', 'Core Tools', 'Manufactura', 'Total', 'Recibieron la pregunta del área', 'Entradas por cada 100 preguntas']); f4 += 1
sem = [d.strftime('%d/%m/%Y') for d in sorted(set(semana(t) for t in E.t_13) | set(semana(t) for t in P.t_pregunta))]
ini4 = f4
for s in sem:
    b = f4 % 2 == 0
    c = celda(w4, f4, 2, s, banda=b); c.alignment = Alignment(horizontal='left')
    for i, e in enumerate(EMB):
        celda(w4, f4, 3 + i, f'=COUNTIFS({ecol("B")},"{e}",{ecol("C")},$B{f4})', ENT, banda=b)
    celda(w4, f4, 6, f'=SUM(C{f4}:E{f4})', ENT, True, MAGENTA, b)
    celda(w4, f4, 7, f'=COUNTIFS({bcol("C")},$B{f4})', ENT, banda=b)
    celda(w4, f4, 8, f'=IF(G{f4}=0,0,F{f4}/G{f4}*100)', DEC, banda=b)
    f4 += 1
celda(w4, f4, 2, 'Total', bold=True)
for c in range(3, 8):
    L = get_column_letter(c); celda(w4, f4, c, f'=SUM({L}{ini4}:{L}{f4 - 1})', ENT, True, MAGENTA if c == 6 else INK)
celda(w4, f4, 8, f'=IF(G{f4}=0,0,F{f4}/G{f4}*100)', DEC, True)
f4 += 1
nota(w4, f4, 'La última semana está incompleta (corte del martes 27/10).'); f4 += 2
seccion(w4, f4, 'De dónde llegan a 1.3', 7); f4 += 1
cab_tabla(w4, f4, ['Venía de', 'ISO + Calidad', 'Core Tools', 'Manufactura', 'Total']); f4 += 1
for d in sorted(set(r[4] for r in filas_e)):
    b = f4 % 2 == 0
    celda(w4, f4, 2, d, banda=b)
    for i, e in enumerate(EMB):
        celda(w4, f4, 3 + i, f'=COUNTIFS({ecol("B")},"{e}",{ecol("E")},$B{f4})', ENT, banda=b)
    celda(w4, f4, 6, f'=SUM(C{f4}:E{f4})', ENT, True, MAGENTA, b)
    f4 += 1

# ================================================================ Funciona
w5 = wb.create_sheet('Funciona 1.3')
encabezado(w5, '¿Funciona 1.3 cuando se usa?', 'Qué pasa con cada entrada a 1.3 y cómo les va a quienes contestaron el área según si los movieron o no', [3, 52, 15, 15, 15, 15, 3])
f5 = 5
seccion(w5, f5, 'Qué pasa dentro de 1.3', 5); f5 += 1
cab_tabla(w5, f5, ['Indicador', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); f5 += 1
E_INI = f5
filas_func = [('Entradas a 1.3', [], None),
              ('El bot mandó "¿te hago una llamadita?"', [('G', 'Sí')], 'base'),
              ('El bot mandó "facilidades de pago"', [('H', 'Sí')], 'base'),
              ('Una persona le escribió mientras estaba en 1.3', [('K', 'Sí')], 'base'),
              ('El cliente respondió en 1.3', [('I', 'Sí')], 'base'),
              ('Se creó una tarea de llamada', [('L', 'Sí')], 'base'),
              ('Siguiente etapa: 1.4 Facilidades de pago', [('M', '1.4 Facilidades de pago')], 'base'),
              ('Siguiente etapa: medio de pago', [('M', 'Medio de pago')], 'base'),
              ('Siguiente etapa: Ganado', [('M', 'Ganado')], 'base'),
              ('Siguiente etapa: Negativo (lo mandó el bot de 1.3)', [('M', 'Negativo'), ('N', 'Bot de 1.3')], 'base'),
              ('Siguiente etapa: Perdido', [('M', 'Perdido')], 'base'),
              ('Sigue en 1.3', [('M', 'Sigue en 1.3')], 'base'),
              ('Después llegó a 1.4 o más (en cualquier momento)', [('P', 'Sí')], 'base'),
              ('Después llegó a un medio de pago', [('Q', 'Sí')], 'base'),
              ('Después Ganado', [('R', 'Sí')], 'base'),
              ('Respondió en 1.3 y aun así el bot lo mandó al Negativo', [('S', 'Sí')], 'base')]
for txt, pares, _ in filas_func:
    b = f5 % 2 == 0
    celda(w5, f5, 2, txt, banda=b)
    celda(w5, f5, 3, f'={cnt(*pares, hoja="e")}' if pares else f'=COUNTA({ecol("A")})', ENT, True, MAGENTA, b)
    for i, e in enumerate(EMB):
        celda(w5, f5, 4 + i, f'={cnt_e(e, *pares, hoja="e")}', ENT, banda=b)
    f5 += 1
f5 += 1
seccion(w5, f5, 'Las mismas cifras en porcentaje de las entradas', 5); f5 += 1
cab_tabla(w5, f5, ['Indicador', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); f5 += 1
for k, (txt, pares, _) in enumerate(filas_func[1:], start=1):
    b = f5 % 2 == 0
    celda(w5, f5, 2, txt, banda=b)
    for c in range(3, 7):
        L = get_column_letter(c)
        celda(w5, f5, c, div(f'{L}{E_INI + k}', f'{L}${E_INI}'), PCT, c == 3, MAGENTA if c == 3 else INK, b)
    f5 += 1
f5 += 1
seccion(w5, f5, 'Respuesta a los mensajes del bot de 1.3', 5); f5 += 1
cab_tabla(w5, f5, ['Mensaje', 'Enviados', 'Respondieron en 30 min', '% respuesta en 30 min']); f5 += 1
BOT13 = {}
for txt, col_env, col_resp in [('"Igual si gustas te puedo hacer una llamadita…"', 'G', 'U'), ('"…tenemos facilidades de pago 💳 ¿Quieres que te muestre cómo quedaría?"', 'H', 'V')]:
    b = f5 % 2 == 0
    BOT13[col_env] = f5
    celda(w5, f5, 2, txt, banda=b)
    celda(w5, f5, 3, f'={cnt((col_env, "Sí"), hoja="e")}', ENT, banda=b)
    celda(w5, f5, 4, f'={cnt((col_resp, "Sí"), hoja="e")}', ENT, banda=b)
    celda(w5, f5, 5, div(f'D{f5}', f'C{f5}'), PCT, True, MAGENTA, b)
    f5 += 1
f5 += 1
seccion(w5, f5, 'Según si el cliente respondió dentro de 1.3', 5); f5 += 1
cab_tabla(w5, f5, ['Grupo', 'Entradas', 'Después 1.4 o más', '% 1.4 o más', 'Después Ganado', '% Ganado']); f5 += 1
RESP13 = {}
for txt, v in [('El cliente respondió en 1.3', 'Sí'), ('El cliente no respondió en 1.3', 'No')]:
    b = f5 % 2 == 0
    RESP13[v] = f5
    celda(w5, f5, 2, txt, banda=b)
    celda(w5, f5, 3, f'={cnt(("I", v), hoja="e")}', ENT, banda=b)
    celda(w5, f5, 4, f'={cnt(("I", v), ("P", "Sí"), hoja="e")}', ENT, banda=b)
    celda(w5, f5, 5, div(f'D{f5}', f'C{f5}'), PCT, True, MAGENTA, b)
    celda(w5, f5, 6, f'={cnt(("I", v), ("R", "Sí"), hoja="e")}', ENT, banda=b)
    celda(w5, f5, 7, div(f'F{f5}', f'C{f5}'), PCT, banda=b)
    f5 += 1
f5 += 1
seccion(w5, f5, 'Tiempo dentro de 1.3 (mediana)', 5); f5 += 1
cab_tabla(w5, f5, ['Indicador', 'Total', 'ISO + Calidad', 'Core Tools', 'Manufactura']); f5 += 1
for txt, col, conds in [('Horas en 1.3', 'O', []), ('Minutos a la respuesta del cliente (si respondió)', 'J', [f'{ecol("I")}="Sí"'])]:
    b = f5 % 2 == 0
    celda(w5, f5, 2, txt, banda=b)
    for i, e in enumerate([None] + EMB):
        ref = f'{get_column_letter(3 + i)}{f5}'
        cc = conds + ([f'{ecol("B")}="{e}"'] if e else [f'{ecol("A")}<>""'])
        w5[ref] = ArrayFormula(ref, '=' + mediana(ecol(col), cc))
        w5[ref].number_format = DEC; w5[ref].font = fnt(10, i == 0, MAGENTA if i == 0 else INK); w5[ref].border = BORDE; w5[ref].alignment = Alignment(horizontal='center')
        if b: w5[ref].fill = fill(SUPERFICIE)
    f5 += 1
f5 += 1
seccion(w5, f5, 'Contestaron el área a tiempo: movidos a 1.3 contra no movidos', 5); f5 += 1
nota(w5, f5, 'Solo leads con 7 días o más desde la pregunta, para que todos hayan tenido el mismo tiempo de seguimiento. Son pocos casos: tomar como referencia, no como efecto demostrado.'); f5 += 1
cab_tabla(w5, f5, ['Grupo', 'Leads', 'Llegaron a 1.4 o más', 'Llegaron a un medio de pago', 'Venta', '% venta']); f5 += 1
GRUPO = {}
for txt, pares in [('Movidos a 1.3', [('N', 'Sí')]), ('No movidos a 1.3', [('N', 'No')]), ('No movidos: el bot los mandó al Negativo', [('N', 'No'), ('S', 'Negativo (lo mandó el bot)')]),
                   ('No movidos: pasaron directo a 1.4, pago o Ganado', [('N', 'No'), ('S', '1.4, pago o Ganado (persona)')]), ('No movidos, con audio', [('N', 'No'), ('K', 'Sí')]), ('No movidos, sin audio', [('N', 'No'), ('K', 'No')])]:
    GRUPO[txt] = f5
    b = f5 % 2 == 0
    base = [('E', 'A tiempo (en 1.2)'), ('X', 'Sí')] + pares
    celda(w5, f5, 2, txt, banda=b)
    celda(w5, f5, 3, f'={cnt(*base)}', ENT, banda=b)
    celda(w5, f5, 4, f'={cnt(*base, ("U", "Sí"))}', ENT, banda=b)
    celda(w5, f5, 5, f'={cnt(*base, ("V", "Sí"))}', ENT, banda=b)
    celda(w5, f5, 6, f'={cnt(*base, ("W", "Sí"))}', ENT, banda=b)
    celda(w5, f5, 7, div(f'F{f5}', f'C{f5}'), PCT, True, MAGENTA, b)
    f5 += 1

# ================================================================ hojas de datos
def hoja_datos(nombre, titulo, sub, cab, filas, anchos, fmts):
    w = wb.create_sheet(nombre)
    w.sheet_view.showGridLines = False
    w['A1'] = titulo; w['A1'].font = Font(name=DISPLAY, size=16, color=MAGENTA)
    w['A2'] = sub; w['A2'].font = fnt(9, color=GRIS)
    w.row_dimensions[1].height = 30
    for i, t in enumerate(cab, start=1):
        c = w.cell(row=4, column=i, value=t)
        c.font = fnt(9, True, 'FFFFFF'); c.fill = fill(MAGENTA); c.border = BORDE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        w.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    w.row_dimensions[4].height = 44
    for rr, f in enumerate(filas, start=5):
        for i, v in enumerate(f, start=1):
            if v is None or v == '': continue
            w.cell(row=rr, column=i, value=v)
    w.freeze_panes = 'B5'
    w.auto_filter.ref = f'A4:{get_column_letter(len(cab))}{len(filas) + 4}'
    return w

wb_base = hoja_datos('Base leads', 'Base de leads', 'Un renglón por lead que recibió la pregunta del área en los 3 embudos. Minutos en tiempo real (no solo horario laboral).',
                     CAB_BASE, filas_base, [8, 20, 11, 12, 20, 10, 11, 26, 13, 12, 11, 26, 10, 10, 24, 12, 16, 22, 26, 11, 11, 11, 13, 10, 24],
                     {3: 'dd/mm/yyyy', 6: DEC, 9: DEC, 13: DEC, 16: DEC})
wb_e = hoja_datos('Entradas a 1.3', 'Entradas a 1.3', 'Un renglón por cada vez que un lead entró a 1.3 Llamadita y facilidades de pago (5/10 al 27/10/2026).',
                  CAB_E, filas_e, [8, 20, 11, 16, 18, 24, 11, 12, 11, 10, 11, 10, 22, 22, 10, 11, 12, 10, 14, 24, 13, 13],
                  {3: 'dd/mm/yyyy', 4: 'dd/mm/yyyy hh:mm', 10: DEC, 15: DEC})
wb_base['A4'].comment = Comment(f'Fuente: CRM SyS (crm-web-production-3b33.up.railway.app), leads de los embudos ISO + Calidad, Core Tools Premium y Manufactura y Calidad, {NB} renglones. Leído por Claude con la cuenta de Joceline (solo lectura), {CORTE_TXT}. Detalle en "Fuentes y método".', 'Claude')
wb_e['A4'].comment = Comment(f'Fuente: CRM SyS, historial de etapas y mensajes, {NE} renglones. Leído por Claude con la cuenta de Joceline (solo lectura), {CORTE_TXT}. Detalle en "Fuentes y método".', 'Claude')

# ================================================================ Fuentes y método
w6 = wb.create_sheet('Fuentes y método')
encabezado(w6, 'Fuentes y método', 'De dónde salen los datos y cómo se definió cada cosa', [3, 34, 110, 3])
f6 = 5
seccion(w6, f6, 'Definiciones', 2); f6 += 1
DEFS = [
    ('Pregunta del área', 'Mensaje del bot de 1.2 Lista de cursos: "Solo una pregunta, ¿en qué área trabajas actualmente? o ¿estás estudiando?". Después el bot espera 2 horas y manda al lead al Negativo.'),
    ('Contestó a tiempo', 'El cliente escribió después de la pregunta y antes de que el lead saliera de 1.2 (normalmente, antes de las 2 horas en que el bot lo manda al Negativo).'),
    ('Contestó tarde', 'Escribió después de la pregunta, pero cuando el lead ya estaba en el Negativo u otra etapa.'),
    ('Audio', 'Nota de voz mandada por una persona del equipo después de la respuesta del cliente. La plantilla "Excelente {{nombre}}! Te mando un audio rapidito…" se cuenta aparte.'),
    ('Movido a 1.3', 'Movimiento a la etapa 1.3 Llamadita y facilidades de pago registrado en el historial del lead después de la pregunta. Todos los movimientos a 1.3 del periodo fueron manuales.'),
    ('Bot de 1.3', 'Al entrar a 1.3: espera 5 min, manda "¿te hago una llamadita?", espera 10 min, manda "facilidades de pago", espera 4 h y manda al lead al Negativo. No se detiene si el cliente contesta.'),
    ('1.4 o más', '1.4 Facilidades de pago, cualquier medio de pago (una sola exhibición, parcialidades, MSI, transferencia, PayPal, tarjeta), STBY o Ganado.'),
    ('Venta', 'Lead en Ganado o con un pago confirmado en el sistema de ventas ligado al cliente, sin cancelar, desde la fecha de la pregunta.'),
    ('Con 7 días o más', 'La pregunta del área se mandó al menos 7 días antes del corte, para comparar grupos con el mismo tiempo de seguimiento.'),
    ('Franja horaria', 'Hora de Reynosa (America/Matamoros) en que el cliente contestó.'),
    ('Minutos', 'Tiempo real de reloj, incluidas noches y fines de semana.'),
]
for a, b2 in DEFS:
    celda(w6, f6, 2, a, bold=True); c = celda(w6, f6, 3, b2, alin='left'); c.alignment = Alignment(wrap_text=True, vertical='center')
    w6.row_dimensions[f6].height = 32
    f6 += 1
f6 += 1
seccion(w6, f6, 'Fuentes de datos', 2); f6 += 1
FUENTES = [
    ('Escrito en', "'Base leads'!A4:Y" + str(B_FIN) + " y 'Entradas a 1.3'!A4:V" + str(E_FIN)),
    ('Fuente', 'CRM SyS, https://crm-web-production-3b33.up.railway.app'),
    ('Objeto', 'GET /embudos, /conversaciones?embudo=…, /conversaciones/:id/movimientos, /conversaciones/:id/mensajes, /conversaciones/:id/tareas, /conversaciones/:id (compras), /etapas/:id/secuencia, /plantillas'),
    ('Cómo se obtuvo', 'Lo corrió Claude con la sesión de Joceline Benavides (rol vendedor, eje Cursos), solo consultas de lectura; no se movió, contestó ni marcó como leída ninguna conversación.'),
    ('Pedido por', 'Rosbel Salinas'),
    ('Parámetros', 'Embudos ISO + Calidad, Core Tools Premium y Manufactura y Calidad; leads cuya pregunta del área o entrada a 1.3 fue del 5/10/2026 en adelante.'),
    ('Leído', CORTE_TXT),
    ('Renglones', f'{NB} leads en Base leads y {NE} entradas en Entradas a 1.3'),
    ('Cambios después de leer', 'Se clasificó cada lead con reglas sobre sus mensajes y su historial de etapas (ver Definiciones). Se dejaron fuera nombres y teléfonos; cada lead se identifica con su número del CRM.'),
    ('Notas', 'Joceline ve los leads del eje Cursos y los que aún no tienen eje: incluye todos los de estos 3 embudos. Un lead que hoy está en otro embudo no entra, aunque haya pasado por estos.'),
]
for a, b2 in FUENTES:
    celda(w6, f6, 2, a, bold=True); c = celda(w6, f6, 3, b2, alin='left'); c.alignment = Alignment(wrap_text=True, vertical='center')
    w6.row_dimensions[f6].height = 30
    f6 += 1

# Logo y colores de pestaña
img = XLImage(LOGO); img.width = img.height = 54
ws.add_image(img, 'F1')
for w, col in [(ws, MAGENTA), (w2, ROSA), (w3, ROSA), (w4, ROSA), (w5, ROSA), (wb_base, INK), (wb_e, INK), (w6, '6B6B6B')]:
    w.sheet_properties.tabColor = col

# ---------------------------------------------------------------- hallazgos (formulas sobre las cifras)
QU, FU = "'Quién lo usa'", "'Funciona 1.3'"
JO = FILA_PERSONA.get('Joceline Sarahi Benavides Serrano')
otros = [f'{QU}!H{r}' for p_, r in FILA_PERSONA.items() if p_ not in ('Joceline Sarahi Benavides Serrano', 'Nadie')]
k = KPI
def T(ref, f='#,##0'):
    return f'TEXT({ref},"{f}")'
HALL = [
    ('¿Se hace el paso a 1.3?', f'="Sí, en la mayoría de los casos, pero no siempre. De "&{T(f"C{k['resp']}")}&" leads que contestaron el área a tiempo, "&{T(f"C{k['audio']}")}&" recibieron audio y "&{T(f"C{k['mov']}")}&" ("&{T(f"C{k['mov_pct']}", "0%")}&") fueron movidos a mano a 1.3."'),
    ('El detalle', f'="Joceline (auxiliar) atendió "&{T(f"{QU}!C{JO}")}&" de esas respuestas y mandó el audio en el "&{T(f"{QU}!G{JO}", "0%")}&", pero solo el "&{T(f"{QU}!H{JO}", "0%")}&" llegó a 1.3, porque otra persona lo movió después. Salma, Ana Victoria y José los mueven entre el "&TEXT(MIN({",".join(otros)}),"0%")&" y el "&TEXT(MAX({",".join(otros)}),"0%")&"."'),
    ('Qué les pasa a los que no se mueven', f'={T(f"C{k['neg_bot']}")}&" leads contestaron a tiempo, no pasaron a 1.3 y a las 2 h el bot les mandó el Negativo. Con 7 días o más, compraron "&{T(f"{FU}!F{GRUPO['No movidos: el bot los mandó al Negativo']}")}&" de "&{T(f"{FU}!C{GRUPO['No movidos: el bot los mandó al Negativo']}")}&" ("&{T(f"{FU}!G{GRUPO['No movidos: el bot los mandó al Negativo']}", "0.0%")}&"), contra "&{T(f"{FU}!G{GRUPO['Movidos a 1.3']}", "0.0%")}&" de los movidos a 1.3. Son pocos casos: es una referencia, no un efecto demostrado."'),
    ('Velocidad', f'="Cuando se hace, se hace rápido: la mediana es de "&{T(f"C{k['min13']}", "0")}&" minutos entre la respuesta del cliente y el paso a 1.3. Todos los movimientos a 1.3 son manuales."'),
    ('¿Funciona 1.3?', f'="Sí, cuando el cliente contesta. Respondió en "&TEXT(IF(C{k['e13']}=0,0,C{k['e13_resp']}/C{k['e13']}),"0%")&" de las "&{T(f"C{k['e13']}")}&" entradas; de esos, el "&{T(f"{FU}!E{RESP13['Sí']}", "0%")}&" llegó a 1.4 o más, contra el "&{T(f"{FU}!E{RESP13['No']}", "0%")}&" de los que no contestaron. A la llamadita responde el "&{T(f"C{k['r30l']}", "0%")}&" en 30 min y a facilidades de pago el "&{T(f"C{k['r30f']}", "0%")}&"."'),
    ('La falla dentro de 1.3', f'="En "&{T(f"C{k['e13_resp_neg']}")}&" entradas el cliente respondió en 1.3 (con un sí, una pregunta de costo o pidiendo la llamada más tarde) y a las 4 h el bot lo mandó al Negativo porque nadie lo pasó a 1.4. El bot de 1.3 no se detiene cuando el cliente contesta."'),
    ('La llamadita', f'="Casi nunca se agenda: solo "&{T(f"C{k['e13_tarea']}")}&" de "&{T(f"C{k['e13']}")}&" entradas a 1.3 tuvieron una tarea de llamada."'),
]
RECO = [
    ('Esta semana', 'Que el paso a 1.3 no dependa de la persona. Lo inmediato: que Joceline mueva a 1.3 al mandar el audio. Lo más seguro: que el CRM pase solo a 1.3 a quien contesta la pregunta del área y avise al asesor.'),
    ('Bot de 1.3', 'Que no mande al Negativo a quien respondió dentro de 1.3: que se detenga cuando el cliente contesta o que, en lugar del Negativo, le cree una tarea al asesor.'),
    ('Prueba "más directo"', 'Corregir primero el paso manual y usar estas cifras de 1.3 como punto de partida (respuesta a la llamadita y a facilidades, avance a 1.4 y venta), para que la prueba mida el mensaje y no la falla del paso a 1.3.'),
]
fh = HALLAZGOS_FILA
for titulo, formula in HALL:
    a = ws.cell(row=fh, column=2, value=titulo); a.font = fnt(10, True, MAGENTA); a.alignment = Alignment(vertical='top', wrap_text=True)
    ws.merge_cells(start_row=fh, start_column=3, end_row=fh, end_column=6)
    c = ws.cell(row=fh, column=3, value=formula); c.font = fnt(10); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[fh].height = 46
    fh += 1
fh += 1
seccion(ws, fh, 'Qué conviene hacer', 5); fh += 1
for titulo, texto in RECO:
    a = ws.cell(row=fh, column=2, value=titulo); a.font = fnt(10, True, ROSA); a.alignment = Alignment(vertical='top', wrap_text=True)
    ws.merge_cells(start_row=fh, start_column=3, end_row=fh, end_column=6)
    c = ws.cell(row=fh, column=3, value=texto); c.font = fnt(10); c.alignment = Alignment(wrap_text=True, vertical='top')
    c.fill = fill(ROSA_TINTE); a.fill = fill(ROSA_TINTE)
    ws.row_dimensions[fh].height = 40
    fh += 1
assert fh <= HALLAZGOS_FILA + 14, fh
ws.column_dimensions['B'].width = 46
for L in 'CDEF':
    ws.column_dimensions[L].width = 17
for rr in range(HALLAZGOS_FILA, HALLAZGOS_FILA + len(HALL)):
    ws.row_dimensions[rr].height = 64
ws.freeze_panes = None

# Alto de renglon para etiquetas largas (Google Sheets no ajusta solo al importar).
import math
for w in [ws, w2, w3, w4, w5, w6]:
    for row in w.iter_rows(min_row=5, max_row=w.max_row):
        for c in row:
            if c.column > 3 or not isinstance(c.value, str) or c.value.startswith('=') or c.coordinate in w.merged_cells:
                continue
            ancho = w.column_dimensions[c.column_letter].width or 10
            lineas = math.ceil(len(c.value) * 1.05 / max(ancho - 2, 1))
            if lineas > 1 and (w.row_dimensions[c.row].height or 15) < 14 * lineas + 6:
                w.row_dimensions[c.row].height = 14 * lineas + 6

import json
json.dump({'KPI': KPI, 'HALLAZGOS_FILA': HALLAZGOS_FILA, 'NB': NB, 'NE': NE}, open(os.path.join(D, 'libro_meta.json'), 'w'))
wb.save(SALIDA)
print('ok', SALIDA, NB, NE, KPI)
