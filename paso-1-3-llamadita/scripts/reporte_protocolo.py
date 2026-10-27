# Prepara las hojas nuevas del reporte "Llamadita contra pregunta abierta":
#   - 'Tras el audio' (datos: un renglon por lead que contesto el area y recibio audio)
#   - 'Respuestas a la llamadita' (datos: las 95 respuestas clasificadas a mano)
#   - 'Llamadita vs pregunta' (resumen con formulas sobre las dos anteriores)
# Escribe un JSON con los valores de cada hoja (para el conector de Google Sheets)
# y un .xlsx de prueba (libro actual + hojas nuevas) para recalcular con LibreOffice.
import json, os, sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

D, LIBRO, SALIDA_JSON, SALIDA_XLSX = sys.argv[1:5]
A = pd.read_pickle(os.path.join(D, 'protocolo.pkl')).sort_values('lead')
CLAS = json.load(open(os.path.join(D, 'clas_llamadita.json')))
SN = lambda v: 'Sí' if bool(v) else 'No'
R1 = lambda v: None if v is None or pd.isna(v) else round(float(v), 1)

# ---------------------------------------------------------------- datos
cab_a = ['Lead', 'Embudo', 'Quién mandó el audio', 'Franja de su respuesta al área', 'Qué recibió después del audio', 'Pregunta abierta ("¿qué te parece mi recomendación…?")',
         'Pasó a 1.3', 'Min. del audio al primer mensaje que le pide respuesta', 'Abrió el chat antes de ese mensaje', 'Respondió en 30 min', 'Respondió en 2 h',
         'Respondió en 24 h', 'Mensajes que recibió antes de responder (2 h)', 'Horas del audio al Negativo', 'Recibió facilidades de pago en 24 h',
         'Llegó a 1.4 o más', 'Llegó a un medio de pago', 'Venta (Ganado o pago confirmado)', 'Con 7 días o más']
CORTO = {'Joceline Sarahi Benavides Serrano': 'Joceline', 'Salma Santos Hernández': 'Salma', 'José Alexander Morales': 'José', 'Ana Victoria de Leon': 'Ana Victoria'}
filas_a = [[int(r.lead), r.embudo, CORTO.get(r.audio_de, r.audio_de), r.franja, r.protocolo, SN(r.con_pregunta), SN(r.con_13), R1(r.min_empujon),
            SN(r.leyo_antes_empujon) if pd.notna(r.min_empujon) else '', SN(r.resp_30m), SN(r.resp_2h), SN(r.resp_24h), int(r.msgs_antes_resp),
            R1(r.horas_a_neg), SN(r.faci_24h), SN(r.llego_14), SN(r.llego_pago), SN(r.venta), SN(r.maduro)] for _, r in A.iterrows()]
cab_b = ['Lead', 'Qué contestó a "¿te hago una llamadita?" (en 30 min, antes del mensaje de facilidades)']
filas_b = [[int(k), v] for k, v in sorted(CLAS.items(), key=lambda x: int(x[0]))]

NA, NB = len(filas_a), len(filas_b)
FA, FB = NA + 4, NB + 4
def a(c): return f"'Tras el audio'!${c}$5:${c}${FA}"
def b(c): return f"'Respuestas a la llamadita'!${c}$5:${c}${FB}"

hoja_a = [['Tras el audio'], [f'Leads de ISO + Calidad, Core Tools Premium y Manufactura y Calidad que contestaron "¿en qué área trabajas?" a tiempo y recibieron audio (5 al 27/10/2026). Un renglón por lead.'], [None], cab_a] + filas_a
hoja_b = [['Respuestas a la llamadita'], ['Respuestas del cliente en los 30 min siguientes al mensaje del bot "Igual si gustas te puedo hacer una llamadita…" (antes del mensaje de facilidades). Clasificadas a mano, una por una.'], [None], cab_b] + filas_b

# ---------------------------------------------------------------- resumen
PROT = ['Pregunta sin 1.3', '1.3 sin pregunta', '1.3 + pregunta']
NOMBRE = {'Pregunta sin 1.3': 'Pregunta abierta, sin 1.3 (Joceline)', '1.3 sin pregunta': '1.3: llamadita y facilidades (Salma, Ana Victoria, José)', '1.3 + pregunta': 'Pregunta abierta + 1.3 (José)'}
S = []  # filas del resumen (columna A vacia, contenido desde B)
arr = []  # (fila, col) de formulas de matriz
def fila(*vals): S.append([None] + list(vals)); return len(S)
def cnt(*pares, extra=''):
    return 'COUNTIFS(' + ','.join(f'{a(c)},"{v}"' for c, v in pares) + extra + ')'
def pct(num, den): return f'=IF({den}=0,0,{num}/{den})'
def med(col, pares):
    cond = '*'.join(f'({a(c)}="{v}")' for c, v in pares)
    return f'=MEDIAN(IF({cond},IF({a(col)}<>"",{a(col)})))'
COLS = 'CDE'

fila('¿La llamadita agrega valor?')
fila('Pregunta abierta sin pasar a 1.3 contra el paso 1.3 · leads que contestaron el área a tiempo y recibieron audio · 5 al 27/10/2026 · CRM leído sin cambios')
fila(None)
fila('VEREDICTO')
VEREDICTO = len(S) + 1
for _ in range(6): fila(None)
fila(None)
fila('LOS TRES CAMINOS DESPUÉS DEL AUDIO')
fila('Indicador', *[NOMBRE[p] for p in PROT]); CAB1 = len(S)
def bloque(titulo, filas_def, base=()):
    rows = {}
    for txt, kind, spec in filas_def:
        r = fila(txt, *([None] * 3)); rows[txt] = r
        for i, p in enumerate(PROT):
            col = COLS[i]
            pares = [('E', p)] + list(base)
            if kind == 'n':
                S[r - 1][2 + i] = '=' + cnt(*pares)
            elif kind == 'pct':
                S[r - 1][2 + i] = pct(cnt(*pares, *[spec]), cnt(*pares))
            elif kind == 'med':
                S[r - 1][2 + i] = med(spec, pares); arr.append((r, 3 + i))
            elif kind == 'avg':
                crit = ','.join(f'{a(c)},"{v}"' for c, v in pares)
                S[r - 1][2 + i] = f'=IF({cnt(*pares)}=0,0,AVERAGEIFS({a(spec)},{crit}))'
    return rows
R_A = bloque(None, [
    ('Leads', 'n', None),
    ('Minutos del audio al mensaje que pide respuesta (mediana)', 'med', 'H'),
    ('Abrieron el chat antes de ese mensaje', 'pct', ('I', 'Sí')),
    ('Mensajes que reciben antes de responder (promedio)', 'avg', 'M'),
    ('Horas del audio al Negativo (mediana)', 'med', 'N'),
    ('Respondieron en 30 min', 'pct', ('J', 'Sí')),
    ('Respondieron en 2 h', 'pct', ('K', 'Sí')),
    ('Respondieron en 24 h', 'pct', ('L', 'Sí')),
    ('Recibieron facilidades de pago en 24 h', 'pct', ('O', 'Sí')),
])
fila(None)
fila('RESULTADO CON 7 DÍAS O MÁS DE SEGUIMIENTO')
fila('Indicador', *[NOMBRE[p] for p in PROT]); CAB2 = len(S)
R_B = bloque(None, [
    ('Leads con 7 días o más', 'n', None),
    ('Llegaron a 1.4 o más', 'pct', ('P', 'Sí')),
    ('Llegaron a un medio de pago', 'pct', ('Q', 'Sí')),
    ('Venta', 'pct', ('R', 'Sí')),
], base=[('S', 'Sí')])
r = fila('Ventas (número)', None, None, None)
for i, p in enumerate(PROT): S[r - 1][2 + i] = '=' + cnt(('E', p), ('S', 'Sí'), ('R', 'Sí'))
R_B['Ventas (número)'] = r
fila('Con pocos casos las diferencias de venta no son concluyentes: ninguna diferencia de esta tabla es estadísticamente significativa (prueba exacta de Fisher, p > 0.1).')
fila(None)
fila('MISMA COMPARACIÓN, SOLO RESPUESTAS DE 9 A 18 H')
fila('Indicador', *[NOMBRE[p] for p in PROT]); CAB3 = len(S)
R_C = bloque(None, [
    ('Leads', 'n', None),
    ('Respondieron en 2 h', 'pct', ('K', 'Sí')),
    ('Respondieron en 24 h', 'pct', ('L', 'Sí')),
    ('Mensajes que reciben antes de responder (promedio)', 'avg', 'M'),
], base=[('D', '9 a 18 h')])
fila(None)
fila('QUÉ CONTESTAN A "¿TE HAGO UNA LLAMADITA?"')
fila('Respuesta', 'Entradas a 1.3', '% de las llamaditas enviadas'); CAB4 = len(S)
r_env = fila('Llamaditas enviadas por el bot', "=COUNTIFS('Entradas a 1.3'!$G$5:$G$360,\"Sí\")", '=IF(C{0}=0,0,C{0}/C{0})')
S[r_env - 1][3] = f'=IF(C{r_env}=0,0,C{r_env}/C{r_env})'
CLASES = ['Acepta la llamada ya', 'Acepta la llamada para otra hora o día', 'No puede o prefiere mensajes', 'Pregunta precio, pagos o el curso', 'Lo va a revisar o agradece']
R_D = {}
for cl in CLASES:
    r = fila(cl, f'=COUNTIFS({b("B")},"{cl}")', None); S[r - 1][3] = f'=IF($C${r_env}=0,0,C{r}/$C${r_env})'; R_D[cl] = r
r = fila('No contestó en esos 30 min', f'=C{r_env}-SUM(C{R_D[CLASES[0]]}:C{R_D[CLASES[-1]]})', None); S[r - 1][3] = f'=IF($C${r_env}=0,0,C{r}/$C${r_env})'; R_D['nada'] = r
r = fila('Tareas de llamada creadas en 1.3', "=COUNTIFS('Entradas a 1.3'!$L$5:$L$360,\"Sí\")", None); S[r - 1][3] = f'=IF($C${r_env}=0,0,C{r}/$C${r_env})'; R_D['tareas'] = r
fila(None)
fila('QUÉ CONVIENE HACER')
RECO = len(S) + 1
for _ in range(5): fila(None)

# Veredicto y recomendaciones: texto con las cifras como formulas.
P_, T_, J_ = 'C', 'D', 'E'  # columnas: pregunta sin 1.3, 1.3 sin pregunta, 1.3 + pregunta
def T(ref, f='0%'): return f'TEXT({ref},"{f}")'
rA, rB, rD = R_A, R_B, R_D
VER = [
    ('Lo que sí se confirma', f'="La llamadita llega antes de que el lead vea el audio: sale a los "&{T(f"{T_}{rA["Minutos del audio al mensaje que pide respuesta (mediana)"]}", "0")}&" min y solo el "&{T(f"{T_}{rA["Abrieron el chat antes de ese mensaje"]}")}&" había abierto el chat. Y casi nadie toma la llamada: de "&{T(f"C{r_env}", "#,##0")}&" invitaciones, "&{T(f"C{rD[CLASES[0]]}+C{rD[CLASES[1]]}", "#,##0")}&" la aceptaron (la mayoría para otra hora o día) y solo hubo "&{T(f"C{rD["tareas"]}", "#,##0")}&" tareas de llamada."'),
    ('Lo que no se sostiene', f'="La pregunta abierta, tal como se usa hoy, no da más tiempo: sale a los "&{T(f"{P_}{rA["Minutos del audio al mensaje que pide respuesta (mediana)"]}", "0.0")}&" min del audio (solo el "&{T(f"{P_}{rA["Abrieron el chat antes de ese mensaje"]}")}&" había abierto el chat) y, como el lead se queda en 1.2, el bot le manda el Negativo a las "&{T(f"{P_}{rA["Horas del audio al Negativo (mediana)"]}", "0.0")}&" h del audio, contra "&{T(f"{T_}{rA["Horas del audio al Negativo (mediana)"]}", "0.0")}&" h en 1.3. Tampoco reciben menos mensajes: en promedio "&{T(f"{P_}{rA["Mensajes que reciben antes de responder (promedio)"]}", "0.0")}&" antes de responder, contra "&{T(f"{T_}{rA["Mensajes que reciben antes de responder (promedio)"]}", "0.0")}&" en 1.3."'),
    ('Resultado', f'="Prácticamente igual. Responden en 2 h el "&{T(f"{P_}{rA["Respondieron en 2 h"]}")}&" con la pregunta y el "&{T(f"{T_}{rA["Respondieron en 2 h"]}")}&" con 1.3; llegan a 1.4 o más el "&{T(f"{P_}{rB["Llegaron a 1.4 o más"]}")}&" y el "&{T(f"{T_}{rB["Llegaron a 1.4 o más"]}")}&"; compran el "&{T(f"{P_}{rB["Venta"]}", "0.0%")}&" y el "&{T(f"{T_}{rB["Venta"]}", "0.0%")}&". La diferencia real: con la pregunta solo el "&{T(f"{P_}{rA["Recibieron facilidades de pago en 24 h"]}")}&" recibe las facilidades de pago en 24 h, contra "&{T(f"{T_}{rA["Recibieron facilidades de pago en 24 h"]}")}&" en 1.3."'),
    ('La mejor combinación', f'="La pregunta abierta dentro de 1.3 (como lo hace José) tiene la respuesta más alta en 30 min: "&{T(f"{J_}{rA["Respondieron en 30 min"]}")}&", contra "&{T(f"{T_}{rA["Respondieron en 30 min"]}")}&" con la llamadita sola. Es una señal, no una diferencia comprobada."'),
    ('¿Agrega valor 1.3?', '="La llamadita automática no agrega valor medible: el mismo resultado se logra con una pregunta abierta. Lo que sí aporta 1.3 es que todos reciban la oferta de facilidades de pago y que el lead no caiga en el Negativo a las 2 h. Conviene conservar el paso, sin la llamada como mensaje automático."'),
    ('Decisión', '="La llamada pasa a ser una alternativa: se ofrece cuando el lead tiene dudas, pide hablar o pone una objeción, o como botón junto a ver los pagos. No como mensaje automático a los 5 minutos."'),
]
RECOS = [
    ('Ya', 'Que ningún lead que contestó el área se quede en 1.2: con 1.2 el Negativo llega a las 2 h de la pregunta. Mientras se ajusta el bot, el lead pasa a 1.3 después del audio y de la pregunta abierta.'),
    ('Bot de 1.3', 'Cambiar el primer mensaje: en lugar de "¿te hago una llamadita?", la pregunta abierta ("¿qué te parece mi recomendación…?"). Las facilidades de pago, cuando conteste que sí o como botón. Que el bot se detenga si el cliente contesta y que nunca mande el Negativo con una conversación abierta.'),
    ('Rediseño v4', 'Coincide con el cierre de la versión 4: "¿Te muestro tus pagos o te llamo 5 minutos?" y el seguimiento de las 3 h "¿Pudiste escuchar el audio?". La llamada queda como una de dos salidas que elige el lead.'),
    ('Entrada automática', 'Que el CRM pase solo a 1.3 a quien contesta la pregunta del área (o cuando se manda la plantilla del audio) y avise al asesor. Así el paso deja de depender de quién atiende.'),
    ('Cómo medirlo', 'Probar la versión nueva de 1.3 contra la actual con leads asignados al azar, unas 5 semanas (unos 300 leads por grupo para ver de 20% a 30% de avance a 1.4). Medir respuesta en 2 h, avance a 1.4 y venta.'),
]
for i, (t, f) in enumerate(VER):
    S[VEREDICTO - 1 + i] = [None, t, f, None, None]
for i, (t, x) in enumerate(RECOS):
    S[RECO - 1 + i] = [None, t, x, None, None]

meta = {'resumen_filas': len(S), 'veredicto': VEREDICTO, 'reco': RECO, 'cabeceras': [CAB1, CAB2, CAB3, CAB4], 'arrays': arr,
        'secciones': [i + 1 for i, row in enumerate(S) if isinstance(row[1], str) and row[1].isupper()], 'na': NA, 'nb': NB}
def limpio(m): return [[None if v is None else v for v in row] for row in m]
S_sheets = [list(row) for row in S]
for (rr, cc) in arr:
    f = S_sheets[rr - 1][cc - 1]
    S_sheets[rr - 1][cc - 1] = '=ARRAYFORMULA(' + f[1:] + ')'
json.dump({'Llamadita vs pregunta': limpio(S_sheets), 'Tras el audio': limpio(hoja_a), 'Respuestas a la llamadita': limpio(hoja_b), 'meta': meta},
          open(SALIDA_JSON, 'w'), ensure_ascii=False)

# xlsx de prueba
wb = load_workbook(LIBRO)
for nombre, matriz in [('Llamadita vs pregunta', S), ('Tras el audio', hoja_a), ('Respuestas a la llamadita', hoja_b)]:
    ws = wb.create_sheet(nombre)
    for i, row in enumerate(matriz, start=1):
        for j, v in enumerate(row, start=1):
            if v is None: continue
            if nombre == 'Llamadita vs pregunta' and (i, j) in set(map(tuple, arr)):
                ref = ws.cell(row=i, column=j).coordinate
                ws[ref] = ArrayFormula(ref, v)
            else:
                ws.cell(row=i, column=j, value=v)
wb.save(SALIDA_XLSX)
print(json.dumps(meta))
